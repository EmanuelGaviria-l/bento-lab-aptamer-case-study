#!/usr/bin/env python3
"""
Build aptamer_subset_v2.csv from the UTexas Sept 2023 spreadsheet.

Keeps the original 7 targets (73 rows after ssDNA/ssRNA + numeric Kd) and
adds further UniProt-able proteins with >=4 clones and >=2 distinct Kds
(plus the original 7-target baseline). Skips cells, small molecules, whole
viruses, and E. coli RNAP (too large for this extra pass).

Does NOT overwrite aptamer_subset.csv (first-submission baseline).
"""

from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

SOURCE = Path("/Users/manny/Downloads/UTexas Aptamer Database dataset_Sept2023.xlsx")
OUTPUT = Path("/Users/manny/Downloads/aptamer_subset_v2.csv")

# Exact Target strings from the spreadsheet (do not merge aliases).
ORIGINAL_TARGETS = [
    "Thrombin, Human",
    "Vascular Endothelial Growth Factor (VEGF)",
    "Basic fibroblast growth factor (bFGF)",
    "Human immunodeficiency virus type 1 Rev (HIV-1 Rev)",
    "Immunoglobulin E (IgE), Human",
    "Autonomously replicating sequence\u2010binding factor 1 (ABF1) protein",
    "Abrin toxin, Abrus precatorius (A.precatorius)",
]

EXPANSION_TARGETS = [
    "Reverse Transcriptase of Type 1 Human Immunodeficiency Virus (HIV-1 RT)",
    "Proteinase K-resistant isoform (PrPSc)",
    "Streptavidin",
    "Unr (Upstream of N-Ras) protein, Human",
    "DNA polymerase (Taq pol), Thermus acquaticus",
    "Lactoferrin",
    "Mammaglobin B (MGB2)",
    "C-terminal ribonuclease domain of colicin E3 (CRD of colicin E3)",
    "HIV-1 aspartyl protease (PR)",
    "Mammalian translation initiation factor 4A (eIF4A)",
    "Soluble Interleukin 2 Receptor \u03b1 (CD25)",
    "Rituximab, anti-CD20 lgG1 antibody",
    "C-terminal region Recombinant human connective tissue growth factor (rhCTGF)",
    # n >= 4 defined proteins
    "Recombinant HA protein from swine IAV H3 cluster IV",
    "RNA-dependent RNA polymerase (3Dpol)",
    "Nonstructural 5B (NS5B) polymerase, hepatitis C virus (HCV) NS5BΔC55 protein",
    "Mucin 1 (MUC1) recombinant protein MUC1-5TR",
    "Mammaglobin A (MGB1)",
    "Leptin protein, Human",
    "Lactate dehydrogenase (LDH), Human",
    "KpnI isoschizomer Acc65I",
    "Estrogen receptor alpha (ER\u03b1)",
    "SARS-CoV-2 Spike glycoprotein",
    "Cytotoxin CNY",
]

TARGETS_V2 = ORIGINAL_TARGETS + EXPANSION_TARGETS

OUT_COLS = [
    "Serial Number",
    "Name of Aptamer",
    "Target",
    "Type of Nucleic Acid",
    "Aptamer Sequence",
    "Sequence Length",
    "GC Content",
    "Kd (nM)",
    "Journal DOI",
    "Citation",
]


def is_numeric_kd(value) -> bool:
    if value is None or value == "":
        return False
    try:
        number = float(value)
    except (TypeError, ValueError):
        return False
    return number == number


def main() -> None:
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows_iter)]
    index = {name: i for i, name in enumerate(header)}

    wanted = set(TARGETS_V2)
    kept: list[dict] = []
    n_start = 0
    n_type = 0
    n_kd = 0

    for row in rows_iter:
        n_start += 1
        nucleic = row[index["Type of Nucleic Acid"]]
        if nucleic not in ("ssDNA", "ssRNA"):
            continue
        n_type += 1
        kd = row[index["Kd (nM)"]]
        if not is_numeric_kd(kd):
            continue
        n_kd += 1
        target = row[index["Target"]]
        if target not in wanted:
            continue
        kept.append({col: row[index[col]] for col in OUT_COLS})

    wb.close()

    kept.sort(key=lambda r: (str(r["Target"]), float(r["Kd (nM)"]), str(r["Serial Number"])))

    with OUTPUT.open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=OUT_COLS)
        writer.writeheader()
        writer.writerows(kept)

    by_target: dict[str, list[float]] = defaultdict(list)
    for row in kept:
        by_target[str(row["Target"])].append(float(row["Kd (nM)"]))

    print(f"Spreadsheet rows:              {n_start}")
    print(f"After ssDNA/ssRNA:             {n_type}")
    print(f"After numeric Kd:              {n_kd}")
    print(f"After v2 target list:          {len(kept)}")
    print(f"Wrote {OUTPUT}")
    print()
    print(f"{'n':>4}  {'uniqKd':>6}  Target")
    for name in TARGETS_V2:
        kds = by_target.get(name, [])
        uniq = len(set(kds))
        tag = "orig" if name in ORIGINAL_TARGETS else "new"
        print(f"{len(kds):4d}  {uniq:6d}  [{tag}] {name}")

    missing = [name for name in TARGETS_V2 if name not in by_target]
    if missing:
        print("\nWARNING — no rows for:")
        for name in missing:
            print(f"  {name!r}")


if __name__ == "__main__":
    main()
